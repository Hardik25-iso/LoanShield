"""Render the report's Results section to Word from the pipeline outputs.

  report/results_template.md  (hand-written; syntax described at the top of the file)
      -> report/results_section.docx

Every number comes from results/ or data/. A value that is not produced yet is
shown as red [key], and [[AUTHOR: ...]] prompts are shown in red italics for the
team to replace with interpretation. Tables and figures are numbered in order of
appearance and @tab:/@fig: references are resolved, since Word has no \\ref.
Headings, body text and captions use Word's built-in styles, so pasting into the
report's own template adopts its formatting.

Run: python -m loanshield.report
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from loanshield.common import ROOT, cfg, get_logger, path, read_json

log = get_logger("report")
TEMPLATE = ROOT / "report" / "results_template.md"
OUT = ROOT / "report" / "results_section.docx"

MODEL_NAMES = {"logreg": "Logistic regression", "rf": "Random forest", "xgb": "XGBoost", "lgbm": "LightGBM",
               "tfidf_lr": "TF-IDF + LR", "tfidf_svm": "TF-IDF + SVM", "muril": "MuRIL"}
GROUP_NAMES = {"metadata": "metadata", "developer": "developer", "permissions": "permissions",
               "listing_text": "listing text", "review_text": "review text"}
MINUS = "−"
RED = RGBColor(0xC0, 0x1F, 0x1F)
TEXT_WIDTH_CM = 16.0  # A4 with 2.5 cm margins


# --- number formatting ---------------------------------------------------------------
def pm(mean, std) -> str:
    return f"{mean:.2f} ± {std:.2f}".replace("-", MINUS)


def num(v, d=2) -> str:
    return f"{v:.{d}f}".replace("-", MINUS)


def signed(v, d=3) -> str:
    if round(v, d) == 0:
        return f"{0:.{d}f}"  # no "-0.00" / "+0.00"
    return f"{v:+.{d}f}".replace("-", MINUS)


def _csv(*parts) -> pd.DataFrame | None:
    f = path(*parts)
    return pd.read_csv(f) if f.exists() else None


def headline() -> str:
    return cfg().get("evaluation", {}).get("headline_subset", "organic")


def model_name(key: str) -> str:
    return MODEL_NAMES.get(key, key)


# --- value collectors (each returns {key: rendered string}) -------------------------
def data_values() -> dict:
    v = {}
    snaps = sorted(path("raw").glob("rbi_dla_2*.csv"))
    if snaps:
        rbi = pd.read_csv(snaps[-1])
        v["data.rbi_snapshot"] = snaps[-1].stem.removeprefix("rbi_dla_")
        v["data.n_rbi_res"] = f"{rbi['re_name'].nunique():,}"
        v["data.n_rbi_pkgs"] = f"{rbi['package_id'].nunique():,}"
        ext = sorted(path("raw").glob("rbi_dla_extract_*.json"))
        v["data.n_rbi_rows"] = f"{len(read_json(ext[-1])['rows']):,}" if ext else f"{len(rbi):,}"
    cands = pd.read_csv(path("interim", "candidates.csv"))
    src = cands.set_index("app_id")["source"]
    v["data.n_candidates"] = f"{len(cands):,}"
    for s, k in [("search", "search"), ("similar", "similar"), ("rbi_list", "rbi")]:
        v[f"data.n_cand_{k}"] = f"{int((cands.source == s).sum()):,}"
    apps = pd.read_csv(path("interim", "apps.csv"))
    v["data.n_kept"] = f"{len(apps):,}"

    labels = pd.read_csv(path("processed", "labels.csv"))
    labels["organic"] = labels["app_id"].map(src).ne("rbi_list")
    labels["dev"] = labels["app_id"].map(apps.set_index("app_id")["developerId"])
    org = labels[labels.organic]
    for name, d in [("", labels), ("organic_", org)]:
        v[f"data.n_{name}verified"] = f"{int((d.label == 'verified').sum()):,}"
        v[f"data.n_{name}unverified"] = f"{int((d.label == 'unverified').sum()):,}"
    v["data.base_rate_all"] = num((labels.label == "unverified").mean(), 3)
    v["data.base_rate_organic"] = num((org.label == "unverified").mean(), 3)
    v["data.n_developers"] = f"{labels.dev.nunique():,}"
    v["data.n_organic_developers"] = f"{org.dev.nunique():,}"
    year = pd.to_datetime(labels["first_seen_date"], errors="coerce").dt.year
    v["data.n_unv_2026"] = str(int(((labels.label == "unverified") & (year == 2026)).sum()))
    for s in ["rbi_match", "seed", "manual"]:
        v[f"data.n_src_{s.removesuffix('_match')}"] = str(int((labels.label_source == s).sum()))

    dec_f = path("interim", "label_decisions.csv")
    if dec_f.exists():
        dec = pd.read_csv(dec_f).set_index("app_id").reindex(labels.app_id)
        auto = labels.label_source.eq("rbi_match").to_numpy()
        lab = labels.label.to_numpy()
        pkg = dec["pkg_match"].fillna(False).astype(bool).to_numpy()
        v["data.n_pkg_match"] = str(int((auto & pkg & (lab == "verified")).sum()))
        v["data.n_name_match"] = str(int((auto & ~pkg & (lab == "verified")).sum()))
        v["data.n_nomatch_unverified"] = str(int((auto & (lab == "unverified")).sum()))

    qa = path("interim", "label_review_A.csv")
    if qa.exists():
        m = pd.read_csv(qa)["label"].value_counts()
        for lab in ["verified", "unverified", "not_lender"]:
            v[f"data.n_manual_{lab}"] = str(int(m.get(lab, 0)))
    au = path("interim", "label_audit.csv")
    if au.exists():
        a = pd.read_csv(au)
        err = int((~a["correct"].astype(bool)).sum())
        v["data.audit_errors"] = str(err)
        v["data.audit_rate"] = f"{100 * err / len(a):.0f}% of {len(a)}"
    sr = path("interim", "seeds_resolved.csv")
    if sr.exists():
        s = pd.read_csv(sr)
        v["data.n_seed_names"] = str(s["name"].str.lower().nunique())
        v["data.n_seed_unresolved"] = str(s.loc[s.resolution == "unresolved", "name"].str.lower().nunique())

    rc = path("processed", "reviews_clean.parquet")
    if rc.exists():
        r = pd.read_parquet(rc, columns=["app_id"])
        v["data.n_reviews_clean"] = f"{int(r.app_id.isin(labels.app_id).sum()):,}"

    st = path("processed", "label_stats.json")
    if st.exists():
        s = read_json(st)
        v["data.n_excluded"] = f"{s['auto_decisions'].get('excluded', 0):,}"
        v["data.n_review_queue"] = str(s["review_queue"])
        m = s.get("manual", {})
        if m.get("kappa") is not None:
            v["data.kappa"] = num(m["kappa"])
        if m.get("agreement") is not None:
            v["data.agreement"] = f"{100 * m['agreement']:.0f}%"
    return v


def model_key(name: str) -> str:
    return "fusion" if name.startswith("fusion") else name


def fusion_name(name: str) -> str:
    inner = re.match(r"fusion\((.+)\)", name)
    return " + ".join(model_name(p) for p in inner.group(1).split("+")) if inner else name


def main_values() -> dict:
    df = _csv("results", "tables", "main_comparison.csv")
    if df is None:
        return {}
    v = {}
    for r in df.itertuples():
        k = f"main.{r.subset}.{model_key(r.model)}"
        for m in ["pr_auc", "roc_auc", "recall_at_5fpr", "precision", "recall", "f1"]:
            v[f"{k}.{m}"] = pm(getattr(r, f"{m}_mean"), getattr(r, f"{m}_std"))
        if r.model.startswith("fusion"):
            v["main.fusion_name"] = fusion_name(r.model)
    meta = path("results", "tables", "_main_meta.json")
    if meta.exists():
        mm = read_json(meta)
        v["main.best_tabular"] = model_name(mm["best_tabular"])
        v["main.best_text"] = model_name(mm["best_text"])
    return v


def temporal_values() -> dict:
    df = _csv("results", "tables", "temporal_holdout.csv")
    if df is None:
        return {}
    d = df[df.subset == headline()].reset_index(drop=True)
    v = {}
    for i, r in d.head(3).iterrows():
        k = f"temporal.{i + 1}"
        v[f"{k}.model"] = model_name(r.model)
        for m in ["pr_auc", "roc_auc", "recall_at_5fpr"]:
            v[f"{k}.{m}"] = num(r[f"{m}_mean"])
    if len(d):
        r = d.iloc[0]
        v["temporal.cutoff"] = str(r.cutoff)
        v["temporal.n_test_pos"] = str(int(r.n_pos_total))
        v["temporal.n_test"] = str(round(r.n_pos_total / r.base_rate))
        v["temporal.base_rate"] = num(r.base_rate, 3)
        from loanshield.models.splits import load_dataset, temporal_split
        ds = load_dataset("features")
        tr, _, _ = temporal_split(ds, pd.Timestamp(r.cutoff))
        v["temporal.n_train"] = str(len(tr))
        v["temporal.n_train_pos"] = str(int(ds.y.iloc[tr].sum()))
    return v


def ablation_values() -> dict:
    df = _csv("results", "tables", "ablation.csv")
    if df is None:
        return {}
    d = df[df.subset == headline()].set_index("dropped_group")
    full = d.loc["none", "pr_auc_mean"]
    v = {"ablation.full": num(full), "ablation.model": model_name(d.loc["none", "model"].rsplit("-", 1)[0])}
    deltas = {g: d.loc[g, "pr_auc_mean"] - full for g in GROUP_NAMES if g in d.index}
    v.update({f"ablation.{g}": signed(x) for g, x in deltas.items()})
    if deltas:
        v["ablation.most_important"] = GROUP_NAMES[min(deltas, key=deltas.get)]
    return v


def early_values() -> dict:
    df = _csv("results", "tables", "early_detection.csv")
    if df is None:
        return {}
    d = df[df.subset == headline()].set_index("features")
    v = {}
    for feat in ["full", "early14"]:
        if feat in d.index:
            for m in ["pr_auc", "roc_auc"]:
                v[f"early.{feat}.{m}"] = pm(d.loc[feat, f"{m}_mean"], d.loc[feat, f"{m}_std"])
    if "full" in d.index:
        v["early.model"] = model_name(d.loc["full", "model"].rsplit("-", 1)[0])
    ef = path("processed", f"features_early{cfg()['splits']['early_days']}.parquet")
    if ef.exists():
        labels = pd.read_csv(path("processed", "labels.csv"), usecols=["app_id"])
        e = pd.read_parquet(ef, columns=["app_id", "has_reviews"])
        v["early.n_with_reviews"] = str(int(e[e.app_id.isin(labels.app_id)].has_reviews.sum()))
    return v


def check_values() -> dict:
    v = {}
    sm = _csv("results", "tables", "imbalance_weights_vs_smote.csv")
    if sm is not None:
        for r in sm[sm.subset == headline()].itertuples():
            m, imb = r.model.rsplit("-", 1)
            v[f"smote.{m}.{imb}"] = pm(r.pr_auc_mean, r.pr_auc_std)
    lk = _csv("results", "tables", "leakage_random_vs_grouped.csv")
    if lk is not None:
        means = {}
        for r in lk[lk.subset == headline()].itertuples():
            m, split = r.model.rsplit("-", 1)
            v[f"leakage.{m}.{split}"] = pm(r.pr_auc_mean, r.pr_auc_std)
            means[(m, split)] = r.pr_auc_mean
        for m in {k[0] for k in means}:
            if (m, "random") in means and (m, "grouped") in means:
                v[f"leakage.{m}.gap"] = signed(means[(m, "random")] - means[(m, "grouped")], 2)
    return v


def shap_values() -> dict:
    v = {}
    card = path("results", "models", "model_card.json")
    # the card must describe the same run as the SHAP tables (demo models outlive result cleanups)
    if card.exists() and path("results", "tables", "shap_importance.csv").exists():
        c = read_json(card)
        v["shap.model"] = model_name(c.get("tabular_model"))
        v["shap.n_apps"] = str(c.get("n_apps"))
    imp = _csv("results", "tables", "shap_importance.csv")
    if imp is not None:
        from loanshield.explain.shap_report import pretty
        for i, f in enumerate(imp["feature"].head(5), 1):
            v[f"shap.top{i}"] = pretty(f)
    terms = _csv("results", "tables", "text_top_terms.csv")
    if terms is not None:
        from loanshield.explain.shap_report import pretty_term

        def q(ts):
            return ", ".join(f"“{pretty_term(t).strip(chr(34))}”" for t in ts)
        v["text.top_unverified"] = q(terms["unverified_term"].head(5))
        v["text.top_verified"] = q(terms["verified_term"].head(5))
    return v


def collect() -> dict:
    v = {}
    for fn in [data_values, main_values, temporal_values, ablation_values, early_values, check_values, shap_values]:
        try:
            v.update(fn())
        except Exception as e:  # noqa: BLE001 - a broken source must not block the rest
            log.warning("%s failed: %s", fn.__name__, e)
    return v


# --- template parsing ------------------------------------------------------------------
@dataclass
class Block:
    kind: str                      # heading1 | heading2 | para | author | table | figure
    text: str = ""
    id: str | None = None
    caption: str = ""
    rows: list[list[str]] = field(default_factory=list)
    images: list[tuple[str, float]] = field(default_factory=list)


def parse(template: str) -> list[Block]:
    template = re.sub(r"<!--.*?-->", "", template, flags=re.S)
    blocks, lines, i = [], template.splitlines(), 0
    para: list[str] = []

    def flush():
        if para:
            text = " ".join(para).strip()
            is_author = text.startswith("[[AUTHOR:") and text.endswith("]]")
            blocks.append(Block("author" if is_author else "para", text))
            para.clear()

    while i < len(lines):
        line = lines[i].rstrip()
        if m := re.match(r"^::: (table|figure) \{#((?:tab|fig):[\w-]+)\}$", line):
            flush()
            b = Block(m.group(1), id=m.group(2))
            i += 1
            while i < len(lines) and lines[i].strip() != ":::":
                body = lines[i].strip()
                if body.startswith("caption:"):
                    b.caption = body.removeprefix("caption:").strip()
                elif body.startswith("image:"):
                    f, w = body.removeprefix("image:").split()
                    b.images.append((f, float(w)))
                elif body.startswith("|"):
                    b.rows.append([c.strip() for c in body.strip("|").split("|")])
                i += 1
            blocks.append(b)
        elif line.startswith("## "):
            flush()
            blocks.append(Block("heading2", line[3:].strip()))
        elif line.startswith("# "):
            flush()
            blocks.append(Block("heading1", line[2:].strip()))
        elif not line.strip():
            flush()
        else:
            para.append(line.strip())
        i += 1
    flush()
    return blocks


def numbering(blocks: list[Block]) -> dict[str, str]:
    refs, n = {}, {"table": 0, "figure": 0}
    for b in blocks:
        if b.kind in n and b.id:
            n[b.kind] += 1
            refs[b.id] = f"{'Table' if b.kind == 'table' else 'Fig.'} {n[b.kind]}"
    return refs


TOKEN = re.compile(r"(<<[\w.]+>>|\[\[AUTHOR:.*?\]\]|\*[^*]+\*|@(?:tab|fig):[\w-]+)")


def segments(text: str, values: dict, refs: dict, missing: list,
             short_missing: bool = False) -> list[tuple[str, str]]:
    """Split text into (style, text) runs: normal | italic | missing | author.

    short_missing renders unfilled values as [TBD] (table cells, where the row and
    column already say what belongs there) instead of [key].
    """
    out = []
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("<<"):
            key = part[2:-2]
            if key in values:
                out.append(("normal", values[key]))
            else:
                missing.append(key)
                out.append(("missing", "[TBD]" if short_missing else f"[{key}]"))
        elif part.startswith("[[AUTHOR:"):
            out.append(("author", "[AUTHOR: " + part[9:-2].strip() + "]"))
        elif part.startswith("*") and part.endswith("*"):
            out.append(("italic", part[1:-1]))
        elif part.startswith("@"):
            out.append(("normal", refs.get(part[1:], part)))
        else:
            out.append(("normal", part))
    return out


# --- Word rendering ----------------------------------------------------------------------
def add_runs(p, segs, size: float | None = None, bold: bool = False) -> None:
    for style, text in segs:
        r = p.add_run(text)
        r.bold = bold or None
        if size:
            r.font.size = Pt(size)
        if style == "italic":
            r.italic = True
        elif style in ("missing", "author"):
            r.font.color.rgb = RED
            r.italic = style == "author"


def _border(cell, edge: str, sz: int) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    el = OxmlElement(f"w:{edge}")
    for k, val in [("val", "single"), ("sz", str(sz)), ("space", "0"), ("color", "000000")]:
        el.set(qn(f"w:{k}"), val)
    borders.append(el)


def _plain_font(style) -> None:
    """Drop theme-font attributes so the explicit font name applies (Word prefers the theme)."""
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is not None:
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{attr}"), None)


def setup(doc) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.5))
    st = doc.styles
    st["Normal"].font.name = "Times New Roman"
    st["Normal"].font.size = Pt(10)
    st["Normal"].paragraph_format.space_after = Pt(6)
    for name, size in [("Heading 1", 12), ("Heading 2", 10)]:
        s = st[name]
        s.font.name, s.font.size, s.font.bold = "Times New Roman", Pt(size), True
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(12), Pt(6)
    cap = st["Caption"]
    cap.font.name, cap.font.size, cap.font.italic, cap.font.bold = "Times New Roman", Pt(9), False, False
    cap.font.color.rgb = RGBColor(0, 0, 0)
    for name in ("Normal", "Heading 1", "Heading 2", "Caption"):
        _plain_font(st[name])


def caption(doc, label: str, segs, before: bool) -> None:
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0 if before else 4)
    p.paragraph_format.space_after = Pt(4 if before else 10)
    p.paragraph_format.keep_with_next = before
    add_runs(p, [("normal", f"{label}. ")], bold=True)
    add_runs(p, segs)


def column_widths(filled, total_cm: float, size_pt: float) -> list[float]:
    """Number columns get exactly the width their values need; text columns share the rest.

    A not-yet-produced value is sized like a typical one ("0.61 ± 0.18"), not by its
    key name, so the layout is the same before and after the numbers arrive.
    """
    char_cm = 0.0185 * size_pt          # average Times New Roman letter width
    digit_cm = 0.016 * size_pt          # digits, spaces and "±" are narrower
    pad_cm = 0.35                        # cell margins

    def length(segs) -> int:
        return sum(11 if style == "missing" else len(t) for style, t in segs)

    def is_text(j) -> bool:
        return any(c.isalpha() for row in filled[1:] for style, t in row[j] if style != "missing" for c in t)

    ncol = len(filled[0])
    head = [max((len(w) for _, t in filled[0][j] for w in t.split()), default=0) * 1.25 * char_cm + pad_cm
            for j in range(ncol)]  # bold header words must not break
    text = [j for j in range(ncol) if is_text(j)]
    body = [max((length(row[j]) for row in filled[1:]), default=0) * (char_cm if j in text else digit_cm) + pad_cm
            for j in range(ncol)]
    word = [max((len(w) for row in filled[1:] for st, t in row[j] if st != "missing" for w in t.split()),
                default=0) * char_cm + pad_cm for j in range(ncol)]  # longest unbreakable word in the body
    widths = [0.0 if j in text else max(body[j], head[j], 1.6) for j in range(ncol)]  # >= 1.6 cm: never wrap a number
    rest = total_cm - sum(widths)
    if text:
        want = {j: max(body[j], head[j]) for j in text}
        floor = {j: max(head[j], word[j]) for j in text}     # never break a word
        if sum(want.values()) <= rest:                       # everything fits: give extra to text
            extra = (rest - sum(want.values())) / len(text)
            for j in text:
                widths[j] = want[j] + extra
        else:                                                # shrink text columns, wrapping long labels
            share = max(rest - sum(floor.values()), 0) / max(sum(want[j] - floor[j] for j in text), 1e-9)
            for j in text:
                widths[j] = floor[j] + (want[j] - floor[j]) * min(share, 1.0)
    scale = total_cm / sum(widths) if sum(widths) > total_cm else 1.0
    return [w * scale for w in widths]


def render_table(doc, b: Block, values, refs, missing) -> None:
    caption(doc, refs[b.id], segments(b.caption, values, refs, missing), before=True)
    rule_before = set()
    rows = []
    for r in b.rows:
        if all(re.fullmatch(r"-{3,}", c) for c in r if c):
            rule_before.add(len(rows))
        else:
            rows.append(r)
    ncol = max(len(r) for r in rows)
    filled = [[segments(r[j] if j < len(r) else "", values, refs, missing, short_missing=True)
               for j in range(ncol)] for r in rows]
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    size = 8 if ncol > 6 else 9
    widths = column_widths(filled, TEXT_WIDTH_CM if ncol > 3 else TEXT_WIDTH_CM * 0.8, size)
    # left-align the first column and any column holding literal (non-placeholder) text
    text_cols = {0} | {j for j in range(ncol) if any(
        j < len(r) and r[j] and not r[j].startswith("<<") and not re.fullmatch(r"[\d−+.,%± ]+", r[j])
        for r in rows[1:])}
    for i, r in enumerate(rows):
        trPr = t.rows[i]._tr.get_or_add_trPr()  # never split a row across pages
        trPr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            trPr.append(OxmlElement("w:tblHeader"))  # repeat header row if the table breaks
        for j in range(ncol):
            cell = t.cell(i, j)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j in text_cols else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = i < len(rows) - 1  # keep the table on one page
            cell.width = Cm(widths[j])
            add_runs(p, filled[i][j], size=size, bold=(i == 0))
            if i == 0:
                _border(cell, "top", 12)
                _border(cell, "bottom", 6)
            if i in rule_before:
                _border(cell, "top", 4)
            if i == len(rows) - 1:
                _border(cell, "bottom", 12)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render_figure(doc, b: Block, values, refs, missing) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    for k, (fname, frac) in enumerate(b.images):
        src = path("results", "figures", fname)
        if src.exists():
            p.add_run().add_picture(str(src), width=Cm(TEXT_WIDTH_CM * frac))
        else:
            missing.append(f"figure:{fname}")
            add_runs(p, [("missing", f"[figure not produced yet: results/figures/{fname}]")])
        if k < len(b.images) - 1:
            p.add_run("  ")
    caption(doc, refs[b.id], segments(b.caption, values, refs, missing), before=False)


def render(template: str, values: dict) -> tuple[Document, list[str]]:
    blocks = parse(template)
    refs = numbering(blocks)
    doc = Document()
    setup(doc)
    missing: list[str] = []
    for b in blocks:
        if b.kind in ("heading1", "heading2"):
            doc.add_heading(b.text, level=1 if b.kind == "heading1" else 2)
        elif b.kind in ("para", "author"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, segments(b.text, values, refs, missing))
        elif b.kind == "table":
            render_table(doc, b, values, refs, missing)
        elif b.kind == "figure":
            render_figure(doc, b, values, refs, missing)
    return doc, missing


def main() -> None:
    template = TEMPLATE.read_text(encoding="utf-8")
    values = collect()
    doc, missing = render(template, values)
    doc.save(OUT)

    keys = set(re.findall(r"<<([\w.]+)>>", re.sub(r"<!--.*?-->", "", template, flags=re.S)))
    miss_keys = sorted({m for m in missing if not m.startswith("figure:")})
    miss_figs = sorted({m.removeprefix("figure:") for m in missing if m.startswith("figure:")})
    groups = pd.Series([k.split(".")[0] for k in miss_keys]).value_counts().to_dict() if miss_keys else {}
    print(f"filled {len(keys) - len(miss_keys)}/{len(keys)} placeholders -> {OUT.relative_to(ROOT)}")
    print(f"missing by section: {groups or 'none'}")
    print(f"figures not produced yet: {miss_figs or 'none'}")
    print(f"author prompts to write: {re.sub(r'<!--.*?-->', '', template, flags=re.S).count('[[AUTHOR:')}")
    (ROOT / "report" / "_placeholders_missing.json").write_text(json.dumps(miss_keys, indent=1), encoding="utf-8")


if __name__ == "__main__":
    main()
